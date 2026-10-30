#!/usr/bin/env python3
"""
ats_check.py — Resume ATS (Applicant Tracking System) friendliness check + keyword coverage

Usage:
    python ats_check.py --resume resume.md --industry internet
    python ats_check.py --resume resume.md --industry tech --jd jd.txt
    python ats_check.py --resume resume.docx --industry finance --out report.md

Supported inputs: .md / .txt / .docx (docx uses python-docx; requires pip install python-docx)
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

INDUSTRIES = {"internet", "tech", "finance", "general"}


def load_resume_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError:
            print(
                "✗ python-docx is missing; install it first: pip install python-docx --break-system-packages",
                file=sys.stderr,
            )
            sys.exit(1)
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    print(f"✗ Unsupported format: {suffix}", file=sys.stderr)
    sys.exit(1)


def load_keywords(industry: str, references_dir: Path) -> list[str]:
    path = references_dir / "keywords" / f"{industry}.txt"
    if not path.exists():
        print(f"✗ Keyword library not found: {path}", file=sys.stderr)
        sys.exit(1)
    return [
        line.strip()
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def coverage(resume_text: str, keywords: list[str]) -> dict:
    text_lower = resume_text.lower()
    hits, missing = [], []
    for kw in keywords:
        # Case-insensitive matching
        if kw.lower() in text_lower:
            hits.append(kw)
        else:
            missing.append(kw)
    return {
        "hits": hits,
        "missing": missing,
        "rate": len(hits) / len(keywords) if keywords else 0,
    }


def jd_extract_keywords(jd_text: str) -> list[str]:
    """Extract candidate keywords from JD text. Simple version: grab common skill / tool / verb tokens."""
    # Grab content words: English tokens up to 30 chars
    candidates = re.findall(
        r"[A-Za-z][A-Za-z0-9+/.\-_]{1,29}",
        jd_text,
    )
    # Simple stop-word removal
    stop = {
        "company", "work", "we", "you", "will", "team", "need", "needs", "able",
        "familiar", "familiarity", "good", "excellent", "experience", "ability",
        "role", "responsibilities", "requirements", "above", "related", "and",
        "including", "perform", "complete", "responsible", "drive", "driving",
        "achieve", "improve", "the", "with", "for", "of", "or", "to", "be", "as",
        "an", "is", "are", "in", "on", "at", "by", "all", "our", "us", "your",
    }
    seen = set()
    out = []
    for c in candidates:
        key = c.lower()
        if key in seen or key in stop:
            continue
        seen.add(key)
        out.append(c)
    return out[:80]  # keep the first 80 candidates


def ats_friendliness(text: str, source_path: Path) -> tuple[int, list[str]]:
    """Assess ATS friendliness; returns (score / 10, warning list)."""
    score = 10
    warnings = []

    # Length
    if len(text) < 200:
        score -= 3
        warnings.append("⚠ Resume text is very short (< 200 characters); parsing may fail or content may be insufficient")
    elif len(text) > 6000:
        score -= 1
        warnings.append("⚠ Resume is on the long side (> 6000 characters); consider trimming to 1-2 pages")

    # Contact info
    has_email = bool(re.search(r"[\w.\-+]+@[\w.\-]+\.\w+", text))
    has_phone = bool(re.search(r"(\+?86[-\s]?)?1[3-9]\d{9}|\d{3}[-\s]?\d{4}[-\s]?\d{4}", text))
    if not has_email:
        score -= 1
        warnings.append("⚠ No email address found")
    if not has_phone:
        score -= 1
        warnings.append("⚠ No phone number found")

    # Decorative symbols
    decorative = re.findall(r"[★☆●○◆◇▶▷■□▪▫♦]", text)
    if len(decorative) > 5:
        score -= 1
        warnings.append(f"⚠ Found {len(decorative)} decorative symbols (★●◆ etc.); ATS may misparse them — trim them")

    # emoji
    emojis = re.findall(r"[\U0001F300-\U0001FAFF\U0001F600-\U0001F64F]", text)
    if emojis:
        score -= 1
        warnings.append(f"⚠ Detected {len(emojis)} emoji; some ATS render them as garbage characters — remove them")

    # Number density (are results sufficiently quantified?)
    numbers = re.findall(r"\d+(?:\.\d+)?%?", text)
    bullet_count = len(re.findall(r"^\s*[-*•]\s+", text, flags=re.MULTILINE))
    if bullet_count > 0:
        density = len(numbers) / bullet_count
        if density < 0.4:
            score -= 1
            warnings.append(
                f"⚠ Quantification density is low (avg {density:.2f} numbers per bullet); "
                f"add more numbers to work / project bullets"
            )

    # docx-specific checks
    if source_path.suffix.lower() == ".docx":
        try:
            from docx import Document
            doc = Document(str(source_path))
            tables = len(doc.tables)
            images = sum(1 for s in doc.inline_shapes)
            if tables > 1:
                score -= 1
                warnings.append(
                    f"⚠ The docx contains {tables} tables; some ATS drop fields when parsing tables — "
                    f"convert them to body paragraphs"
                )
            if images > 0:
                score -= 1
                warnings.append(
                    f"⚠ The docx contains {images} images (including the ID photo); ATS don't read images — "
                    f"don't put important info only in images; the photo may stay"
                )
        except Exception:
            pass

    return max(0, score), warnings


def render_report(
    industry_cov: dict,
    industry: str,
    jd_cov: dict | None,
    ats_score: int,
    ats_warnings: list[str],
) -> str:
    lines = [
        "# Resume ATS Check Report",
        "",
        f"## Industry Keyword Coverage ({industry})",
        f"- Hit rate: **{industry_cov['rate'] * 100:.1f}%** "
        f"({len(industry_cov['hits'])} / {len(industry_cov['hits']) + len(industry_cov['missing'])})",
        "",
        "**Covered**: " + (", ".join(industry_cov["hits"]) or "(none)"),
        "",
        "**Suggested additions** (first 15): " + (", ".join(industry_cov["missing"][:15]) or "(none)"),
        "",
    ]

    if jd_cov is not None:
        lines += [
            "## JD Keyword Coverage",
            f"- Hit rate: **{jd_cov['rate'] * 100:.1f}%**",
            "",
            "**Covered**: " + (", ".join(jd_cov["hits"][:30]) or "(none)"),
            "",
            "**In the JD but missing from the resume** (prioritize these): " + (", ".join(jd_cov["missing"][:20]) or "(none)"),
            "",
        ]

    lines += [
        f"## ATS Friendliness Score: **{ats_score}/10**",
        "",
    ]
    if ats_warnings:
        lines += ats_warnings
    else:
        lines.append("✅ No obvious problems")

    lines += [
        "",
        "---",
        "## Improvement Priority",
        "",
        "1. First close the JD hit-rate gap → the most direct signal for ATS pass rate",
        "2. Then add industry keywords → so the resume surfaces in broader searches",
        "3. Finally tune ATS friendliness → remove decorative symbols, emoji, tables, images",
        "",
        "Note: higher coverage isn't always better — **keywords must appear inside real achievement bullets**.",
        "Don't list a long keyword string as your skills section; HR will see through it.",
    ]
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--resume", required=True, help="Resume file path (.md/.txt/.docx)")
    parser.add_argument(
        "--industry", choices=list(INDUSTRIES), required=True, help="Industry direction"
    )
    parser.add_argument("--jd", help="Optional: JD text file for precise comparison")
    parser.add_argument("--out", help="Output report path; prints to stdout by default")
    parser.add_argument(
        "--references-dir",
        default=str(Path(__file__).resolve().parent.parent / "references"),
    )
    args = parser.parse_args()

    resume_path = Path(args.resume).expanduser()
    if not resume_path.exists():
        print(f"✗ Resume file not found: {resume_path}", file=sys.stderr)
        sys.exit(1)

    text = load_resume_text(resume_path)
    references_dir = Path(args.references_dir)

    industry_keywords = load_keywords(args.industry, references_dir)
    industry_cov = coverage(text, industry_keywords)

    jd_cov = None
    if args.jd:
        jd_path = Path(args.jd).expanduser()
        if not jd_path.exists():
            print(f"✗ JD file not found: {jd_path}", file=sys.stderr)
            sys.exit(1)
        jd_keywords = jd_extract_keywords(jd_path.read_text(encoding="utf-8"))
        jd_cov = coverage(text, jd_keywords)

    ats_score, ats_warnings = ats_friendliness(text, resume_path)
    report = render_report(industry_cov, args.industry, jd_cov, ats_score, ats_warnings)

    if args.out:
        out_path = Path(args.out).expanduser()
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(report, encoding="utf-8")
        print(f"✓ Report written to: {out_path}")
    else:
        print(report)


if __name__ == "__main__":
    main()
